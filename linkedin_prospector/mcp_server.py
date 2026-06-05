#!/usr/bin/env python3
"""
Enterprise Sales Prospector — MCP Server

Multi-source contact and company intelligence platform.

Tools:
  search_prospects          : LinkedIn people search with Claude scoring (US filter support)
  scrape_linkedin_profile   : Extract full profile data from a LinkedIn URL
  scrape_linkedin_profiles  : Bulk scrape multiple LinkedIn profile URLs
  parse_html_prospects      : Score profiles from saved HTML (LinkedIn fallback)
  find_linkedin_url         : Find exact LinkedIn profile URL via Google research
  bulk_find_linkedin_urls   : Batch LinkedIn URL resolution for a list of names
  search_indian_company     : Search Indian companies on Tofler + Zaubacorp
  get_company_directors     : Extract directors from Indian company registry
  google_search_people      : Free-form Google search for professional profiles
  export_to_excel           : Export a list of contacts to a styled Excel file
  list_us_locations         : List supported US city/state location filters
  cache_stats               : Show cache statistics and storage savings
"""

import asyncio
import json
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(__file__))

from mcp.server.fastmcp import FastMCP

mcp = FastMCP("enterprise-prospector")


# ---------------------------------------------------------------------------
# Tool: search_prospects  (LinkedIn people search + Claude scoring)
# ---------------------------------------------------------------------------

@mcp.tool()
async def search_prospects(
    keywords: str = "",
    designation: str = "",
    process: str = "",
    location: str = "",
    us_only: bool = False,
    company: str = "",
    pages: int = 3,
    min_score: int = 0,
    export_excel: bool = True,
    auth_state_path: str = "auth_state.json",
    anthropic_api_key: str = "",
) -> str:
    """
    Search LinkedIn for prospects and score them with Claude AI.

    Goes directly to linkedin.com/search/results/people and extracts profiles.
    Requires auth_state.json (run: playwright codegen --save-storage=auth_state.json https://www.linkedin.com).
    Results are cached to avoid repeat scraping and API costs.

    Args:
        keywords: Domain keywords, e.g. "AI automation SaaS"
        designation: Target job title, e.g. "VP of Sales" or "CPO"
        process: Business process, e.g. "procurement" or "digital transformation"
        location: Location filter — city, state, or country name (e.g. "New York", "California")
        us_only: If true, restricts results to United States (overrides location if set)
        company: Company name filter, e.g. "Airbus" or "TCS"
        pages: LinkedIn result pages to scrape (1-5, default 3)
        min_score: Only return profiles with score >= this value (0-10)
        export_excel: If true, save results to an Excel file
        auth_state_path: Path to saved Playwright auth state (login cookies)
        anthropic_api_key: Anthropic API key (falls back to ANTHROPIC_API_KEY env)
    """
    from sources.linkedin_source import scrape_people_search
    from scorer import score_profiles

    if not any([keywords, designation, process, location, company]):
        return json.dumps({"error": "Provide at least one search criterion."})

    profiles_raw = await scrape_people_search(
        keywords=keywords or "",
        designation=designation or "",
        process=process or "",
        location=location or "",
        company=company or "",
        us_only=us_only,
        max_pages=max(1, min(pages, 5)),
        auth_state_path=auth_state_path,
    )

    if not profiles_raw:
        return json.dumps({
            "error": (
                "No profiles found. Ensure auth_state.json is valid "
                "or use parse_html_prospects as fallback."
            ),
            "profiles": [],
        })

    profiles = [p.to_dict() for p in profiles_raw]

    api_key = anthropic_api_key or os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return json.dumps({
            "message": "No Anthropic API key — returning unscored profiles.",
            "total": len(profiles),
            "profiles": profiles,
        })

    scored = score_profiles(
        profiles=profiles,
        keywords=keywords or None,
        designation=designation or None,
        process=process or None,
        location=location or None,
        company=company or None,
        api_key=api_key,
    )

    if min_score > 0:
        scored = [p for p in scored if p.get("score", 0) >= min_score]
    scored.sort(key=lambda p: p.get("score", 0), reverse=True)

    if export_excel:
        path = _export_contacts_excel(scored, f"prospects_{date.today()}")
        for p in scored:
            p["_exported_to"] = path

    return json.dumps({"total": len(scored), "profiles": scored}, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: scrape_linkedin_profile  (full profile data extraction)
# ---------------------------------------------------------------------------

@mcp.tool()
async def scrape_linkedin_profile(
    linkedin_url: str,
    auth_state_path: str = "auth_state.json",
) -> str:
    """
    Go to a LinkedIn profile URL and extract all visible data.

    Extracts: name, headline, location, about/summary, full experience history,
    education, skills, and contact info (if visible).
    Results are cached for 24 hours.

    Args:
        linkedin_url: Full LinkedIn profile URL, e.g. https://www.linkedin.com/in/satyanadella/
        auth_state_path: Path to saved Playwright auth state (required for full data)
    """
    from sources.linkedin_source import scrape_profile

    result = await scrape_profile(linkedin_url, auth_state_path)
    return json.dumps(result, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: scrape_linkedin_profiles  (bulk profile extraction)
# ---------------------------------------------------------------------------

@mcp.tool()
async def scrape_linkedin_profiles(
    urls_json: str,
    auth_state_path: str = "auth_state.json",
    concurrency: int = 2,
    export_excel: bool = True,
) -> str:
    """
    Bulk scrape multiple LinkedIn profiles and extract full data from each.

    Visits each profile URL on linkedin.com and extracts name, headline,
    location, experience, education, skills, and contact info.
    Runs with controlled concurrency to avoid rate limiting.

    Args:
        urls_json: JSON array of LinkedIn profile URLs
                   Example: ["https://www.linkedin.com/in/satyanadella/", ...]
        auth_state_path: Path to saved Playwright auth state
        concurrency: Parallel profile requests (1-3, default 2)
        export_excel: Save all profiles to an Excel file
    """
    from sources.linkedin_source import scrape_profiles_bulk

    try:
        urls = json.loads(urls_json)
    except json.JSONDecodeError as e:
        return json.dumps({"error": f"Invalid JSON: {e}"})

    if not isinstance(urls, list):
        return json.dumps({"error": "urls_json must be a JSON array of strings."})

    concurrency = max(1, min(concurrency, 3))
    profiles = await scrape_profiles_bulk(urls, auth_state_path, concurrency)

    if export_excel and profiles:
        path = _export_contacts_excel(profiles, f"linkedin_profiles_{date.today()}", "LinkedIn Profiles")
        for p in profiles:
            if isinstance(p, dict):
                p["_exported_to"] = path

    return json.dumps({
        "total": len(profiles),
        "profiles": profiles,
    }, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: parse_html_prospects  (LinkedIn HTML fallback)
# ---------------------------------------------------------------------------

@mcp.tool()
async def parse_html_prospects(
    html: str,
    keywords: str = "",
    designation: str = "",
    process: str = "",
    location: str = "",
    company: str = "",
    min_score: int = 0,
    export_excel: bool = True,
    anthropic_api_key: str = "",
) -> str:
    """
    Parse LinkedIn profiles from raw HTML (fallback when scraping is blocked).

    Save a LinkedIn search results page as HTML from your browser
    (File → Save Page As) and paste the content here.

    Args:
        html: Raw HTML content of a LinkedIn search results page
        keywords / designation / process / location / company: Scoring criteria
        min_score: Only return profiles with score >= this value (0-10)
        export_excel: Save results to Excel file
        anthropic_api_key: Anthropic API key (falls back to ANTHROPIC_API_KEY env)
    """
    from scraper import scrape_from_html
    from scorer import score_profiles

    profiles = scrape_from_html(html)
    if not profiles:
        return json.dumps({"error": "No profiles found in the provided HTML."})

    api_key = anthropic_api_key or os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return json.dumps({"message": "No API key — raw profiles returned.", "profiles": profiles})

    scored = score_profiles(
        profiles=profiles,
        keywords=keywords or None,
        designation=designation or None,
        process=process or None,
        location=location or None,
        company=company or None,
        api_key=api_key,
    )

    if min_score > 0:
        scored = [p for p in scored if p.get("score", 0) >= min_score]
    scored.sort(key=lambda p: p.get("score", 0), reverse=True)

    if export_excel:
        path = _export_contacts_excel(scored, f"prospects_{date.today()}")
        for p in scored:
            p["_exported_to"] = path

    return json.dumps({"total": len(scored), "profiles": scored}, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: find_linkedin_url  (single person lookup via Google)
# ---------------------------------------------------------------------------

@mcp.tool()
async def find_linkedin_url(
    name: str,
    company: str = "",
    title: str = "",
    location: str = "",
) -> str:
    """
    Find the exact LinkedIn profile URL for a person using Google search.

    Does NOT consume LinkedIn sessions or ZoomInfo credits.
    Uses Google's index and caches results for 48 hours.

    Args:
        name: Full name of the person, e.g. "Satya Nadella"
        company: Current or recent company, e.g. "Microsoft"
        title: Job title for disambiguation, e.g. "CEO"
        location: City or country for disambiguation
    """
    from sources.google_source import find_linkedin_url as _find

    url = await _find(name=name, company=company, title=title, location=location)
    return json.dumps({
        "name": name,
        "company": company,
        "linkedin_url": url,
        "found": bool(url),
    })


# ---------------------------------------------------------------------------
# Tool: bulk_find_linkedin_urls  (batch URL resolution)
# ---------------------------------------------------------------------------

@mcp.tool()
async def bulk_find_linkedin_urls(
    contacts_json: str,
    concurrency: int = 3,
) -> str:
    """
    Resolve exact LinkedIn profile URLs for a batch of contacts via Google.

    Zero LinkedIn sessions consumed — uses Google search + caching.
    Much cheaper than ZoomInfo enrichment for URL-only lookups.

    Args:
        contacts_json: JSON array with keys: name, company, title (optional)
                       Example: [{"name": "Ratan Tata", "company": "Tata Sons"}]
        concurrency: Parallel Google requests (1-5, default 3)
    """
    from sources.google_source import find_linkedin_url as _find

    try:
        contacts = json.loads(contacts_json)
    except json.JSONDecodeError as e:
        return json.dumps({"error": f"Invalid JSON: {e}"})

    concurrency = max(1, min(concurrency, 5))
    sem = asyncio.Semaphore(concurrency)

    async def _resolve(c: dict) -> dict:
        async with sem:
            url = await _find(
                name=c.get("name", ""),
                company=c.get("company", ""),
                title=c.get("title", ""),
                location=c.get("location", ""),
            )
            return {**c, "linkedin_url": url, "found": bool(url)}

    results = await asyncio.gather(*[_resolve(c) for c in contacts])
    found = sum(1 for r in results if r["found"])
    return json.dumps({
        "total": len(results),
        "found": found,
        "not_found": len(results) - found,
        "results": list(results),
    }, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: search_indian_company
# ---------------------------------------------------------------------------

@mcp.tool()
async def search_indian_company(
    company_name: str,
    source: str = "both",
) -> str:
    """
    Search for an Indian company on Tofler and/or Zaubacorp (MCA registry).

    Returns CIN, status, incorporation date, registered address, directors.
    Ideal for finding Indian subsidiaries, unlisted entities, or joint ventures.

    Args:
        company_name: Company name or partial name to search
        source: "tofler", "zaubacorp", or "both" (default)
    """
    results = []

    if source in ("both", "tofler"):
        from sources.tofler_source import search_company as tofler_search
        r = await tofler_search(company_name)
        results.extend([{**x, "source": "tofler"} for x in r])

    if source in ("both", "zaubacorp"):
        from sources.zaubacorp_source import search_company as zauba_search
        r = await zauba_search(company_name)
        results.extend([{**x, "source": "zaubacorp"} for x in r])

    seen: set[str] = set()
    deduped = []
    for r in results:
        cin = r.get("cin", "")
        if cin and cin in seen:
            continue
        if cin:
            seen.add(cin)
        deduped.append(r)

    return json.dumps({
        "query": company_name,
        "total": len(deduped),
        "results": deduped,
    }, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: get_company_directors
# ---------------------------------------------------------------------------

@mcp.tool()
async def get_company_directors(
    company_name: str = "",
    cin: str = "",
    source: str = "zaubacorp",
    find_linkedin: bool = False,
) -> str:
    """
    Extract directors and key people from an Indian company registry.

    Pulls director names, DINs, and designations from MCA data.
    Optionally resolves LinkedIn profiles for each director via Google.

    Args:
        company_name: Company name to search (used when CIN is not known)
        cin: Corporate Identity Number for direct lookup (faster, more accurate)
        source: "tofler" or "zaubacorp" (default)
        find_linkedin: If true, resolve LinkedIn URLs for each director via Google
    """
    from utils.normalizer import Contact

    contacts: list[Contact] = []

    if source == "tofler":
        from sources.tofler_source import get_directors_as_contacts, get_company_details
        if cin:
            company = await get_company_details(f"https://www.tofler.in/company/{cin}")
            for d in company.directors:
                parts = d["name"].split()
                contacts.append(Contact(
                    full_name=d["name"],
                    first_name=parts[0] if parts else "",
                    last_name=" ".join(parts[1:]) if len(parts) > 1 else "",
                    designation=d.get("designation", "Director"),
                    company=company.name,
                    source="tofler",
                    confidence=0.9,
                ))
        elif company_name:
            contacts = await get_directors_as_contacts(company_name)
    else:
        from sources.zaubacorp_source import get_directors_as_contacts, get_company_by_cin
        if cin:
            company = await get_company_by_cin(cin)
            for d in company.directors:
                parts = d["name"].split()
                contacts.append(Contact(
                    full_name=d["name"],
                    first_name=parts[0] if parts else "",
                    last_name=" ".join(parts[1:]) if len(parts) > 1 else "",
                    designation=d.get("designation", "Director"),
                    company=company.name,
                    source="zaubacorp",
                    confidence=0.9,
                ))
        elif company_name:
            contacts = await get_directors_as_contacts(company_name)

    results = [c.to_dict() for c in contacts]

    if find_linkedin and results:
        from sources.google_source import find_linkedin_url
        sem = asyncio.Semaphore(3)

        async def _enrich(c: dict) -> dict:
            async with sem:
                url = await find_linkedin_url(
                    name=c["full_name"],
                    company=c.get("company", ""),
                    title=c.get("designation", ""),
                )
                c["linkedin_url"] = url
                return c

        results = list(await asyncio.gather(*[_enrich(c) for c in results]))

    return json.dumps({
        "source": source,
        "total": len(results),
        "directors": results,
    }, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: google_search_people
# ---------------------------------------------------------------------------

@mcp.tool()
async def google_search_people(
    query: str,
    max_results: int = 10,
) -> str:
    """
    Free-form Google search returning professional profiles and LinkedIn URLs.

    Use this for broad discovery with no source constraints.
    Examples:
      "CTO Fintech New York 2024"
      "Partner McKinsey digital transformation"
      "site:linkedin.com/in VP Engineering San Francisco"

    Args:
        query: Google search query string
        max_results: Max results to return (default 10, max 20)
    """
    from sources.google_source import google_search_contacts

    results = await google_search_contacts(
        query=query,
        max_results=min(max_results, 20),
    )
    return json.dumps({
        "query": query,
        "total": len(results),
        "results": results,
    }, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool: list_us_locations
# ---------------------------------------------------------------------------

@mcp.tool()
async def list_us_locations() -> str:
    """
    List all supported US city, state, and metro area location filters.

    Use these names in the 'location' parameter of search_prospects.
    LinkedIn geo URNs are resolved automatically.
    """
    from sources.linkedin_source import US_GEO_URNS

    grouped = {"cities_metros": [], "states": []}
    state_keywords = [
        "california", "texas", "florida", "new york state", "illinois",
        "pennsylvania", "ohio", "georgia", "north carolina", "michigan",
        "new jersey", "virginia", "washington state", "massachusetts",
    ]
    for name in sorted(US_GEO_URNS.keys()):
        if name in ("us", "usa", "united states"):
            continue
        if name in state_keywords:
            grouped["states"].append(name)
        else:
            grouped["cities_metros"].append(name)

    return json.dumps({
        "tip": "Pass any of these names as the 'location' parameter in search_prospects.",
        "us_only_shortcut": "Set us_only=true to restrict to all of United States without specifying a city.",
        **grouped,
    }, indent=2)


# ---------------------------------------------------------------------------
# Tool: export_to_excel
# ---------------------------------------------------------------------------

@mcp.tool()
async def export_to_excel(
    contacts_json: str,
    filename: str = "",
    sheet_title: str = "Prospects",
) -> str:
    """
    Export a list of contacts to a professionally formatted Excel file.

    Accepts any JSON array of contact objects. Automatically detects fields,
    creates styled headers, alternating row colors, and clickable hyperlinks.

    Args:
        contacts_json: JSON array of contact objects
        filename: Output filename (without .xlsx). Defaults to prospects_YYYY-MM-DD
        sheet_title: Excel sheet name (max 31 chars)
    """
    try:
        contacts = json.loads(contacts_json)
    except json.JSONDecodeError as e:
        return json.dumps({"error": f"Invalid JSON: {e}"})

    fname = filename or f"prospects_{date.today()}"
    path = _export_contacts_excel(contacts, fname, sheet_title)
    return json.dumps({"exported_to": path, "rows": len(contacts)})


# ---------------------------------------------------------------------------
# Tool: cache_stats
# ---------------------------------------------------------------------------

@mcp.tool()
async def cache_stats() -> str:
    """
    Show cache statistics — entries stored per source and total size.

    The cache prevents re-scraping the same pages and saves ZoomInfo credits.
    Cached results expire after 24–168 hours depending on data type.
    """
    from utils.cache import Cache
    stats = Cache().stats()
    return json.dumps(stats, indent=2)


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _export_contacts_excel(
    contacts: list[dict],
    filename: str,
    sheet_title: str = "Prospects",
) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "openpyxl not installed — run: pip install openpyxl"

    if not contacts:
        return ""

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    NAVY, WHITE, GREEN, GRAY = "1F3864", "FFFFFF", "C6EFCE", "F2F2F2"

    # Collect all field keys, skip internal _ fields
    all_keys: list[str] = []
    for c in contacts:
        for k in c:
            if k not in all_keys and not k.startswith("_") and k != "raw":
                all_keys.append(k)

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    hdr_font = Font(color=WHITE, bold=True, size=11, name="Calibri")
    green_fill = PatternFill("solid", fgColor=GREEN)
    gray_fill = PatternFill("solid", fgColor=GRAY)

    for col, key in enumerate(all_keys, 1):
        cell = ws.cell(row=1, column=col, value=key.replace("_", " ").title())
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(key) + 4)
    ws.row_dimensions[1].height = 28

    url_cols = {i + 1 for i, k in enumerate(all_keys)
                if any(x in k.lower() for x in ["url", "linkedin", "link", "website"])}

    for row_idx, contact in enumerate(contacts, 2):
        row_fill = green_fill if row_idx % 2 == 0 else gray_fill
        for col, key in enumerate(all_keys, 1):
            val = contact.get(key, "")
            # Flatten lists/dicts for Excel
            if isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.font = Font(size=10, name="Calibri")
            if col in url_cols and isinstance(val, str) and val.startswith("http"):
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single", size=10, name="Calibri")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(all_keys))}{len(contacts) + 1}"

    path = f"{filename}.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------

if __name__ == "__main__":
    mcp.run()
