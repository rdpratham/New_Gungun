import os
import io
import re
from datetime import datetime
from flask import Flask, request, jsonify, render_template, send_file
from flask_cors import CORS
import anthropic
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__)
CORS(app)

ANTHROPIC_API_KEY = os.environ.get('ANTHROPIC_API_KEY', '')

MOM_PROMPT = """You are an expert corporate meeting minutes writer. Generate a detailed, professional Minutes of Meeting (MOM) document from the transcript below.

METADATA PROVIDED (use exactly as given):
- Date: {date}
- Time: {time}
- Mode: {mode}
- Subject: {subject}

INSTRUCTIONS:
1. Extract from the transcript: organization names, all attendees with their roles/titles, duration, and a "Prepared By" organization.
2. Write numbered discussion points — each as a full, detailed paragraph capturing the substance of what was said. Be thorough.
3. Extract all action items with owners and timelines into a structured table.
4. Derive strategic next steps.
5. If any field cannot be determined from the transcript, use a reasonable placeholder like "As per transcript" or "TBD".

OUTPUT FORMAT — follow this structure EXACTLY (use plain text, no markdown asterisks or hashes):

MINUTES OF MEETING
[Org 1]  |  [Org 2] — [Short Description of Meeting]

Date\t{date}
Time\t{time}
Mode\t{mode}
Duration\t[Estimated or stated duration]
Prepared By\t[Preparing organization]
Attendees ([Org 1])\t[Full names with roles]
Attendees ([Org 2])\t[Full names with roles]
Subject\t{subject}

DISCUSSION SUMMARY
The following is a structured account of key points discussed during the meeting.
1.  [Detailed paragraph — first discussion point]
2.  [Detailed paragraph — second discussion point]
[Continue for all significant points discussed...]

ACTION ITEMS
Agreed actions following this discussion.
Action Item\tOwner\tTimeline
[Action 1]\t[Owner]\t[Timeline]
[Action 2]\t[Owner]\t[Timeline]
[Continue for all action items...]

NEXT STEPS
Strategic priorities to advance the engagement.
1.  [Next step 1]
2.  [Next step 2]
[Continue for all next steps...]

This document is confidential and intended solely for the parties named herein.

TRANSCRIPT:
{transcript}"""

SUMMARY_PROMPT = """Based on the Minutes of Meeting below, write a concise executive call summary.

Structure the summary as follows:
- Opening paragraph: purpose and participants of the meeting (2-3 sentences)
- Key Discussion Points: 3-5 bullet points capturing the most important topics
- Decisions Made: bullet points of any decisions or agreements reached
- Critical Actions: the 3-5 most important action items with owners
- Closing: one sentence on the overall outcome and next milestone

Keep it professional, tight, and under 400 words.

MINUTES OF MEETING:
{mom}"""


def parse_uploaded_file(file):
    filename = file.filename.lower()
    if filename.endswith('.txt'):
        return file.read().decode('utf-8')
    elif filename.endswith('.docx'):
        doc = Document(file)
        return '\n'.join(para.text for para in doc.paragraphs)
    else:
        raise ValueError('Unsupported file format. Please upload a .txt or .docx file.')


def call_claude(prompt):
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    message = client.messages.create(
        model='claude-sonnet-4-6',
        max_tokens=4096,
        messages=[{'role': 'user', 'content': prompt}]
    )
    return message.content[0].text


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/generate-mom', methods=['POST'])
def generate_mom():
    try:
        date = request.form.get('date', '').strip()
        time_val = request.form.get('time', '').strip()
        subject = request.form.get('subject', '').strip()
        mode = request.form.get('mode', '').strip()
        transcript = request.form.get('transcript', '').strip()

        if 'file' in request.files and request.files['file'].filename:
            try:
                transcript = parse_uploaded_file(request.files['file'])
            except ValueError as e:
                return jsonify({'error': str(e)}), 400

        if not transcript:
            return jsonify({'error': 'Please provide a transcript — either paste it or upload a file.'}), 400
        if not date:
            return jsonify({'error': 'Date is required.'}), 400
        if not subject:
            return jsonify({'error': 'Subject is required.'}), 400

        if not ANTHROPIC_API_KEY:
            return jsonify({'error': 'ANTHROPIC_API_KEY is not configured on the server.'}), 500

        prompt = MOM_PROMPT.format(
            date=date,
            time=time_val or 'Not specified',
            mode=mode or 'Not specified',
            subject=subject,
            transcript=transcript
        )

        mom_content = call_claude(prompt)
        return jsonify({'mom': mom_content, 'success': True})

    except anthropic.AuthenticationError:
        return jsonify({'error': 'Invalid Anthropic API key. Please check server configuration.'}), 500
    except Exception as e:
        return jsonify({'error': f'Generation failed: {str(e)}'}), 500


@app.route('/generate-summary', methods=['POST'])
def generate_summary():
    try:
        mom_content = request.json.get('mom', '').strip()
        if not mom_content:
            return jsonify({'error': 'No MOM content provided.'}), 400

        prompt = SUMMARY_PROMPT.format(mom=mom_content)
        summary = call_claude(prompt)
        return jsonify({'summary': summary, 'success': True})

    except Exception as e:
        return jsonify({'error': f'Summary generation failed: {str(e)}'}), 500


def build_docx(content, title='Document'):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue

        # Main title line
        if line.strip() == 'MINUTES OF MEETING' or (i == 0 and 'MINUTES' in line):
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Subtitle (org | org line)
        if i == 1 and ('|' in line or '—' in line):
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x4A, 0x4A, 0x6A)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Section headers
        if line.strip() in ('DISCUSSION SUMMARY', 'ACTION ITEMS', 'NEXT STEPS') or (
            line.strip().isupper() and len(line.strip()) > 3 and '\t' not in line
        ):
            doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
            i += 1
            continue

        # Tab-separated metadata rows (Date\tValue, Time\tValue, etc.)
        if '\t' in line and not line.strip().startswith(('1.', '2.', '3.', '4.', '5.',
                                                          '6.', '7.', '8.', '9.')):
            parts = line.split('\t', 1)
            if len(parts) == 2:
                p = doc.add_paragraph()
                label_run = p.add_run(parts[0].strip() + '\t')
                label_run.bold = True
                label_run.font.size = Pt(10)
                value_run = p.add_run(parts[1].strip())
                value_run.font.size = Pt(10)
                i += 1
                continue

        # Action items table header
        if line.strip() == 'Action Item\tOwner\tTimeline':
            # Collect table rows
            table_rows = []
            j = i + 1
            while j < len(lines) and lines[j].strip() and '\t' in lines[j]:
                table_rows.append(lines[j].split('\t'))
                j += 1

            table = doc.add_table(rows=1 + len(table_rows), cols=3)
            table.style = 'Table Grid'

            # Header row
            hdr_cells = table.rows[0].cells
            for idx, header in enumerate(['Action Item', 'Owner', 'Timeline']):
                hdr_cells[idx].text = header
                for para in hdr_cells[idx].paragraphs:
                    for run in para.runs:
                        run.bold = True
                        run.font.size = Pt(9)

            for row_idx, row_data in enumerate(table_rows):
                row_cells = table.rows[row_idx + 1].cells
                for col_idx in range(3):
                    val = row_data[col_idx].strip() if col_idx < len(row_data) else ''
                    row_cells[col_idx].text = val
                    for para in row_cells[col_idx].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)

            i = j
            continue

        # Numbered discussion points and next steps
        if re.match(r'^\d+\.\s', line.strip()):
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Confidentiality / footer lines
        if 'confidential' in line.lower() or '©' in line:
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            i += 1
            continue

        # Default paragraph
        p = doc.add_paragraph()
        run = p.add_run(line.strip())
        run.font.size = Pt(10)
        i += 1

    return doc


@app.route('/download-docx', methods=['POST'])
def download_docx():
    try:
        content = request.json.get('content', '')
        doc_type = request.json.get('type', 'mom')
        title = 'Minutes of Meeting' if doc_type == 'mom' else 'Call Summary'

        doc = build_docx(content, title)
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"MOM_{timestamp}.docx" if doc_type == 'mom' else f"Summary_{timestamp}.docx"

        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/download-pdf', methods=['POST'])
def download_pdf():
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import inch, cm
        from reportlab.lib import colors

        content = request.json.get('content', '')
        doc_type = request.json.get('type', 'mom')

        buffer = io.BytesIO()
        doc_pdf = SimpleDocTemplate(
            buffer, pagesize=A4,
            rightMargin=1.2 * inch, leftMargin=1.2 * inch,
            topMargin=1 * inch, bottomMargin=1 * inch
        )

        styles = getSampleStyleSheet()

        style_title = ParagraphStyle('MOMTitle', fontSize=16, fontName='Helvetica-Bold',
                                     textColor=colors.HexColor('#1A1A2E'), alignment=TA_CENTER,
                                     spaceAfter=4)
        style_subtitle = ParagraphStyle('MOMSubtitle', fontSize=11, fontName='Helvetica-Bold',
                                        textColor=colors.HexColor('#4A4A6A'), alignment=TA_CENTER,
                                        spaceAfter=12)
        style_section = ParagraphStyle('MOMSection', fontSize=12, fontName='Helvetica-Bold',
                                       textColor=colors.HexColor('#1A1A2E'), spaceBefore=12,
                                       spaceAfter=6)
        style_meta_label = ParagraphStyle('MOMMetaLabel', fontSize=10, fontName='Helvetica-Bold',
                                          spaceAfter=2)
        style_body = ParagraphStyle('MOMBody', fontSize=10, fontName='Helvetica',
                                    leading=14, spaceAfter=6)
        style_footer = ParagraphStyle('MOMFooter', fontSize=8, fontName='Helvetica-Oblique',
                                      textColor=colors.HexColor('#888888'), spaceBefore=12)

        story = []
        lines = content.split('\n')
        i = 0
        table_data = []
        collecting_table = False

        while i < len(lines):
            line = lines[i]

            if not line.strip():
                story.append(Spacer(1, 8))
                i += 1
                continue

            if line.strip() == 'MINUTES OF MEETING':
                story.append(Paragraph(line.strip(), style_title))
                i += 1
                continue

            if i == 1 and ('|' in line or '—' in line):
                story.append(Paragraph(line.strip(), style_subtitle))
                i += 1
                continue

            if line.strip() in ('DISCUSSION SUMMARY', 'ACTION ITEMS', 'NEXT STEPS') or (
                line.strip().isupper() and len(line.strip()) > 3 and '\t' not in line
            ):
                story.append(Paragraph(line.strip(), style_section))
                i += 1
                continue

            # Table header detection
            if line.strip() == 'Action Item\tOwner\tTimeline':
                table_data = [['Action Item', 'Owner', 'Timeline']]
                j = i + 1
                while j < len(lines) and lines[j].strip() and '\t' in lines[j]:
                    row = lines[j].split('\t')
                    table_data.append([r.strip() for r in row[:3]])
                    j += 1

                col_widths = [3.5 * inch, 1.8 * inch, 1.2 * inch]
                t = Table(table_data, colWidths=col_widths)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A1A2E')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1),
                     [colors.HexColor('#F8F8FC'), colors.white]),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
                i = j
                continue

            # Metadata tab rows
            if '\t' in line and not re.match(r'^\d+\.\s', line.strip()):
                parts = line.split('\t', 1)
                if len(parts) == 2:
                    text = f'<b>{parts[0].strip()}</b>    {parts[1].strip()}'
                    story.append(Paragraph(text, style_body))
                    i += 1
                    continue

            if 'confidential' in line.lower() or '©' in line:
                story.append(Paragraph(line.strip(), style_footer))
                i += 1
                continue

            story.append(Paragraph(line.strip(), style_body))
            i += 1

        doc_pdf.build(story)
        buffer.seek(0)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"MOM_{timestamp}.pdf" if doc_type == 'mom' else f"Summary_{timestamp}.pdf"

        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype='application/pdf'
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
